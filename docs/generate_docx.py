"""Word 문서 생성 스크립트."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(6)

    # 표지
    doc.add_paragraph()
    title = doc.add_heading("Homeostatic Financial Agent", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("2층 항상성 구조 기반 투자 행동 출현 실험")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14)
    doc.add_paragraph()
    p = doc.add_paragraph("프로젝트 기술 문서")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("2026-04-03")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # 1. 프로젝트 개요
    doc.add_heading("1. 프로젝트 개요", level=1)

    doc.add_heading("1.1 핵심 아이디어", level=2)
    doc.add_paragraph(
        '"돈을 잃지 않으려는 인간의 마음"을 항상성(homeostasis) 유지 회로로 모델링한다. '
        "에이전트에게 투자하라는 지시를 하지 않고, 구매력(purchasing power)을 일정 수준으로 "
        "유지하라는 항상성 목표만 부여했을 때, 투자 행동이 자발적으로 출현(emergence)하는지 관찰한다."
    )

    doc.add_heading("1.2 이론적 배경", level=2)
    doc.add_paragraph(
        "Maturana & Varela (1984) - autopoiesis(자기생산). "
        "생명체는 자기 자신을 유지하려는 조직 그 자체이며, "
        "인식은 이 자기유지 과정에서 출현한다.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Yoshida et al. (2024) - homeostatic reinforcement learning. "
        "체온, 혈당 등 내부 상태의 항상성 유지만을 목표로 학습시켰더니, "
        "걷기, 먹이찾기, 체온조절 등 복합 행동이 출현함을 확인.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Damasio - somatic marker hypothesis. "
        "감정(신체 상태의 변화)이 의사결정을 편향시키며, 이는 항상성 유지와 직결됨.",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Selten (1998) - aspiration adaptation theory. "
        "성공/실패의 기준선(aspiration level)이 경험에 따라 적응적으로 변화함.",
        style="List Bullet",
    )

    doc.add_heading("1.3 기존 연구와의 차별점", level=2)
    doc.add_paragraph(
        "기존 Agent-Based Model(ABM)은 loss aversion, herding, momentum 등의 행동 규칙을 "
        "파라미터로 직접 주입한다. 본 프로젝트는 행동을 주입하지 않고 "
        "항상성 유지라는 단일 목표에서 행동이 출현하는지 관찰하는 점에서 근본적으로 다르다."
    )

    # 2. 환경 설계
    doc.add_heading("2. 환경 설계 (single_agent_env.py)", level=1)

    doc.add_heading("2.1 2층 항상성 구조", level=2)
    doc.add_paragraph(
        "에이전트의 내부 동기를 2층으로 분리한다. 이는 뇌과학적 근거가 있다:"
    )
    doc.add_paragraph(
        "1층 = 시상하부(hypothalamus), 뇌간: 체온, 혈당 등 생리적 항상성",
        style="List Bullet",
    )
    doc.add_paragraph(
        "2층 = 전전두엽(prefrontal cortex), 전대상피질: 사회적 비교, 자기 평가",
        style="List Bullet",
    )

    doc.add_heading("1층: 생존 항상성 (Survival Homeostasis)", level=3)
    doc.add_paragraph("대상: 구매력의 절대값", style="List Bullet")
    doc.add_paragraph(
        "방향: 양방향 (구매력이 setpoint보다 높아도, 낮아도 스트레스)", style="List Bullet"
    )
    doc.add_paragraph(
        "생물학적 대응: 체온이 36.5도에서 벗어나면 위아래 모두 위험한 것과 동일",
        style="List Bullet",
    )
    doc.add_paragraph("setpoint: 1.0 (고정)", style="List Bullet")

    doc.add_heading("2층: 사회적 항상성 (Social Homeostasis)", level=3)
    doc.add_paragraph(
        "대상: 시장 평균 구매력 대비 나의 상대적 위치 (social_position = my_pp / market_avg_pp)",
        style="List Bullet",
    )
    doc.add_paragraph("방향: 비대칭", style="List Bullet")
    doc.add_paragraph(
        "  상승(남들보다 올라감) = 약한 쾌감 (social_gain_bonus = 0.3)"
    )
    doc.add_paragraph(
        "  하락(남들보다 뒤처짐) = 강한 고통 (social_loss_penalty = 2.0)"
    )
    doc.add_paragraph(
        "근거: 사회적 지위 하락은 뇌에서 신체적 고통과 동일한 회로를 활성화함(전대상피질). "
        "반면 성공을 싫어하는 사람은 없으므로 상승은 스트레스가 아닌 쾌감."
    )

    doc.add_heading("계층적 구조 (Hierarchical)", level=3)
    doc.add_paragraph(
        "Maslow의 욕구 계층과 동일한 원리. 배고프면 자존심 없다. "
        "이 계층 구조는 주입이 아니라 생물학적 사실이다."
    )
    doc.add_paragraph(
        "survival_stress > survival_threshold(0.15)일 때: 1층이 지배, reward = -survival_stress"
    )
    doc.add_paragraph(
        "survival_stress <= threshold, social_deviation >= 0일 때: 2층 지배, reward = +0.3 * social_deviation"
    )
    doc.add_paragraph(
        "survival_stress <= threshold, social_deviation < 0일 때: 2층 지배, reward = +2.0 * social_deviation (음수)"
    )

    doc.add_heading("2.2 기초대사 (Metabolism)", level=2)
    doc.add_paragraph(
        "매 timestep마다 구매력이 일정 비율(metabolism_rate = 0.0002, 일별 기준) 감소한다. "
        "이는 인플레이션과 생활비를 모델링한 것으로, 아무것도 하지 않으면 구매력이 서서히 줄어든다. "
        "생물학적으로 기초대사량에 대응한다."
    )

    doc.add_heading("2.3 시장 평균 구매력", level=2)
    doc.add_paragraph(
        "2층 사회적 항상성을 위해 남들의 평균 구매력을 시뮬레이션한다. "
        "매 step, 시장 평균 구매력도 독립적인 GBM 수익률로 업데이트된다. "
        "에이전트가 투자하지 않아도 남들은 투자하고 있으므로, "
        "가만히 있으면 상대적 위치가 자동으로 하락한다."
    )

    doc.add_heading("2.4 구매력 관측 지연 (Observation Lag)", level=2)
    doc.add_paragraph(
        "현실 투자자는 자신의 실질 구매력을 실시간으로 알지 못한다. "
        "이를 모델링하기 위해, 에이전트가 관측하는 구매력에 n-step 지연(lag)과 노이즈를 추가한다."
    )

    doc.add_heading("2.5 Observation 공간", level=2)
    doc.add_paragraph(
        "2층 + hvol 활성 시 7차원: "
        "[관측된 구매력, survival_deviation, social_position, social_deviation, "
        "최근 자산 수익률, 이전 투자 비율, hvol/100]"
    )

    doc.add_heading("2.6 행동 공간", level=2)
    doc.add_paragraph("위험자산 투자 비율 [0, 1] (continuous).")
    doc.add_paragraph("0 = 전액 현금 보유, 1 = 전액 위험자산 투자")

    doc.add_heading("2.7 가격 프로세스", level=2)
    doc.add_paragraph(
        "Geometric Brownian Motion (GBM): asset_return ~ N(mu=0.0003, sigma=0.015) 일별 기준. "
        "연환산 기대수익률 약 7.5%, 변동성 약 24%. "
        "의도적으로 가장 단순한 가격 프로세스를 사용하여 환경 효과와 항상성 효과를 구분."
    )

    doc.add_heading("2.8 종료 조건", level=2)
    doc.add_paragraph("구매력 <= 0.1: 사망 (terminated), 패널티 -10", style="List Bullet")
    doc.add_paragraph("current_step >= max_steps: 시간 초과 (truncated)", style="List Bullet")

    doc.add_heading("2.9 Historical Volatility (hvol)", level=2)
    doc.add_paragraph(
        "선택적 활성화(enable_hvol=True). 과거 20일간 수익률의 표준편차를 연환산하여 계산. "
        "주의: 이것은 실현 변동성(historical volatility)이지, VIX(내재 변동성, implied volatility)가 아니다."
    )

    # 3. 학습 알고리즘
    doc.add_heading("3. 학습 알고리즘", level=1)

    doc.add_heading("3.1 PPO (Proximal Policy Optimization)", level=2)
    add_table(
        doc,
        ["파라미터", "값"],
        [
            ["구현", "Stable-Baselines3"],
            ["Policy", "MlpPolicy (2-layer FC)"],
            ["learning_rate", "3e-4"],
            ["n_steps", "2048"],
            ["batch_size", "64"],
            ["n_epochs", "10"],
            ["gamma", "0.99"],
            ["gae_lambda", "0.95"],
            ["clip_range", "0.2"],
            ["total_timesteps", "200,000 ~ 300,000"],
        ],
    )

    doc.add_heading("3.2 핵심 설계 원칙", level=2)
    doc.add_paragraph(
        "에이전트는 투자하라는 지시를 받지 않는다. "
        "reward는 오직 항상성 유지 여부에만 의존한다. "
        "투자 행동은 기초대사(구매력 감소 압력) 때문에 항상성을 유지하려면 "
        "수익을 올려야 한다는 것을 에이전트가 스스로 학습해야 한다."
    )

    # 4. 실험 및 결과
    doc.add_heading("4. 실험 및 결과", level=1)

    doc.add_heading("4.1 Phase 1: 단일 항상성 회로 (2026-03-21)", level=2)
    doc.add_paragraph("스크립트: sim/run_experiment.py")
    doc.add_paragraph("1층(생존 항상성)만 활성화한 상태에서 에이전트 학습.")
    doc.add_paragraph(
        "투자 행동 출현 확인. 비선형 반응 함수 출현: deviation이 커질수록 투자 비율 급증."
    )

    doc.add_heading("4.2 Observation Lag 실험", level=2)
    doc.add_paragraph("스크립트: sim/run_lag_comparison.py")
    add_table(
        doc,
        ["lag", "행동 패턴", "해석", "Kurtosis"],
        [
            ["0", "적극적, 안정적", "트레이더", "낮음"],
            ["5", "과잉반응, 불안정", "일반 투자자", "8.72 (fat tail)"],
            ["20", "행동 포기", "예금자", "낮음"],
        ],
    )
    doc.add_paragraph(
        "GBM 환경 자체는 정규분포인데 에이전트의 행동이 fat tail을 보인다. "
        "관측 지연으로 인한 과잉 반응(overshoot)이 원인."
    )

    doc.add_heading("4.3 Phase 1.5: 2층 항상성 실험 (2026-04-02)", level=2)
    doc.add_paragraph("스크립트: sim/run_dual_homeostasis.py")
    add_table(
        doc,
        ["지표", "1층만 (생존)", "2층 (생존 + 사회적)"],
        [
            ["평균 투자 비율", "0.14 (소극적)", "0.68 (적극적)"],
            ["최종 구매력", "0.89 (하락)", "1.18 (상승)"],
            ["Action kurtosis", "3.49 (fat tail)", "-0.99 (안정적)"],
            ["사회적 층 활성 비율", "-", "43.3%"],
        ],
    )
    doc.add_paragraph(
        "사회적 항상성이 투자 적극성을 출현시킴. "
        "1층만 있으면 setpoint 위에서 투자를 중단하지만, 2층을 넣으면 남들 대비 위치 유지를 위해 투자 지속."
    )

    doc.add_heading("4.4 Contrarian 실험 (2026-04-02)", level=2)
    doc.add_paragraph("스크립트: sim/run_contrarian.py")
    add_table(
        doc,
        ["모델", "Original PP", "Contrarian PP", "비고"],
        [
            ["1층", "0.89", "0.96", "수익 증가, 변동성 폭발 (std 0.43)"],
            ["2층", "1.18", "0.93", "오히려 손해"],
        ],
    )
    doc.add_paragraph(
        '"공포에 사라"의 구조적 재현. '
        "1층(공포 회로) 반대 = 수익+고변동, 2층 반대 = 손해."
    )

    doc.add_heading("4.5 실제 시장 데이터 평가 (2026-04-02)", level=2)
    doc.add_paragraph("스크립트: sim/run_real_data.py")
    doc.add_paragraph(
        "S&P 500 10년, 2,513일. 연 13.25% 수익, 변동성 18.05%, kurtosis 16.02."
    )
    add_table(
        doc,
        ["모델", "최종 구매력", "투자 비율", "시장 평균"],
        [
            ["1층", "0.90", "7.7%", "-"],
            ["2층", "1.14", "70.4%", "1.25"],
        ],
    )
    doc.add_paragraph(
        "GBM으로 학습한 에이전트가 실제 시장에서도 비슷하게 작동. 항상성 회로가 환경 변화에 강건."
    )

    doc.add_heading("4.6 LightGBM 예측 실험 (2026-04-02~03)", level=2)
    doc.add_paragraph("스크립트: sim/run_lgbm.py")
    doc.add_paragraph(
        "에이전트 행동(투자 비율)을 피쳐로 LightGBM에 투입. "
        "1층 투자비율 = 공포 신호, 2층 투자비율 = 사회적 압력 신호."
    )
    doc.add_paragraph(
        "데이터: S&P 500 + VIX (2010~2026). 학습: ~2019, 테스트: 2020~."
    )
    add_table(
        doc,
        ["실험", "타겟", "AUC", "비고"],
        [
            ["1", "다음날 수익률 방향", "0.484", "예측력 없음"],
            ["2", "60일 수익률 방향", "0.614", "2층 피쳐 중요도 1등"],
            ["3", "60일 VRP 변화 방향", "0.569", "슬라이딩 윈도우, 3965샘플"],
            ["4", "60일 VRP (60일 에이전트)", "0.569", "60일 단위 재학습"],
        ],
    )
    doc.add_paragraph(
        "종합: 단기 예측 불가, 중기 약한 시그널 존재 (AUC 0.57~0.61). "
        "모든 실험에서 2층 피쳐가 중요도 1등. "
        "한계: GBM 학습 에이전트 vs 실제 시장 도메인 불일치."
    )

    # 5. 파일 구조
    doc.add_heading("5. 파일 구조", level=1)
    add_table(
        doc,
        ["파일", "설명"],
        [
            ["env/single_agent_env.py", "Gymnasium 환경 (2층 항상성, lag, hvol)"],
            ["sim/run_experiment.py", "Phase 1: 단일 에이전트 학습/평가"],
            ["sim/run_lag_comparison.py", "lag 비교 실험 (lag=0,5,20)"],
            ["sim/run_dual_homeostasis.py", "Phase 1.5: 1층 vs 2층 비교"],
            ["sim/run_contrarian.py", "Contrarian 역지표 실험"],
            ["sim/run_real_data.py", "S&P 500 실제 데이터 평가"],
            ["sim/run_lgbm.py", "LightGBM 피쳐 기반 예측"],
            ["analysis/stylized_facts.py", "Stylized facts 분석 도구"],
        ],
    )

    # 6. 파라미터
    doc.add_heading("6. 환경 파라미터 전체 목록", level=1)
    add_table(
        doc,
        ["파라미터", "기본값", "설명"],
        [
            ["max_steps", "1000", "에피소드 최대 길이"],
            ["metabolism_rate", "0.0002", "매 step 구매력 감소율"],
            ["asset_mu", "0.0003", "위험자산 기대수익률 (일별)"],
            ["asset_sigma", "0.015", "위험자산 변동성 (일별)"],
            ["survival_setpoint", "1.0", "1층 항상성 목표 구매력"],
            ["initial_purchasing_power", "1.0", "초기 구매력"],
            ["death_threshold", "0.1", "사망 임계값"],
            ["survival_threshold", "0.15", "1층/2층 전환 경계"],
            ["enable_social", "True", "2층 사회적 항상성 활성화"],
            ["social_setpoint", "1.0", "초기 사회적 위치"],
            ["social_gain_bonus", "0.3", "사회적 상승 보상 강도"],
            ["social_loss_penalty", "2.0", "사회적 하락 패널티 강도"],
            ["observation_lag", "0", "구매력 관측 지연 (step 수)"],
            ["observation_noise", "0.0", "관측 노이즈 표준편차"],
            ["enable_hvol", "False", "실현 변동성 관측 활성화"],
            ["hvol_window", "20", "실현 변동성 계산 윈도우"],
        ],
    )

    # 7. 핵심 발견
    doc.add_heading("7. 핵심 발견 요약", level=1)
    findings = [
        "항상성 회로만으로 투자 행동이 출현한다. 에이전트에게 투자를 가르치지 않았지만, 기초대사를 상쇄하기 위해 자발적으로 투자를 시작한다.",
        "관측 지연(lag)이 투자자 유형을 분화시킨다. lag=0은 트레이더, lag=5는 과잉반응 일반 투자자(fat tail), lag=20은 예금자.",
        "사회적 항상성이 투자 적극성의 원천이다. 1층만 있으면 14%만 투자, 2층 추가 시 68%까지 투자.",
        '"공포에 사라"가 구조적으로 재현된다. 1층 반대 = 수익+고변동, 2층 반대 = 손해.',
        "시장 예측은 약하지만, 사회적 항상성 신호가 가장 유의미하다. 모든 실험에서 2층 피쳐가 중요도 1등.",
        "GBM으로 학습한 에이전트가 실제 시장에서도 비슷하게 작동한다. 항상성 회로가 환경 변화에 강건.",
    ]
    for i, f in enumerate(findings, 1):
        doc.add_paragraph(f"{i}. {f}")

    # 8. 한계
    doc.add_heading("8. 한계 및 향후 과제", level=1)
    doc.add_heading("한계", level=2)
    doc.add_paragraph("GBM 환경에서 학습한 에이전트는 실제 시장의 fat tail, volatility clustering, 레짐 전환을 경험하지 못함", style="List Bullet")
    doc.add_paragraph("단일 위험자산만 존재 (다자산 포트폴리오 미지원)", style="List Bullet")
    doc.add_paragraph("사회적 항상성의 비대칭 파라미터는 외부에서 설정 (출현이 아닌 주입 요소)", style="List Bullet")

    doc.add_heading("향후 과제", level=2)
    doc.add_paragraph("실제 시장 데이터로 에이전트 학습 (도메인 일치)", style="List Bullet")
    doc.add_paragraph("피쳐 확장: 에이전트 행동 외 추가 시장 지표", style="List Bullet")
    doc.add_paragraph("다중 에이전트에서 내생적 가격 형성 및 stylized facts 출현 여부 검증", style="List Bullet")
    doc.add_paragraph("사회적 항상성의 비대칭 파라미터 감도 분석", style="List Bullet")

    doc.save("docs/project_documentation.docx")
    print("Saved: docs/project_documentation.docx")


if __name__ == "__main__":
    main()
